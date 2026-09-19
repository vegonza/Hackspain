import unittest
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from documents.classification import PaymentRules
from erp import ErpEntry
from pipeline.extraction_3.extraction import InvoiceExtraction, InvoiceLine
from documents.payment_notes import PaymentConcern, PaymentNotesReview, SourcedPaymentConcern, SourcedPaymentNotesReview, review_payment_notes


class PaymentNotesTests(unittest.TestCase):
    def test_rejects_evidence_not_present_in_invoice(self) -> None:
        review = SourcedPaymentNotesReview(concerns=[SourcedPaymentConcern(note_index=1, reason="Anulación")])
        with patch("documents.payment_notes.create_extractor") as factory:
            factory.return_value.__enter__.return_value.run.return_value = review
            with self.assertRaisesRegex(ValueError, "evidence is absent"):
                review_payment_notes(["Condiciones de pago: 30 días."], {})

    def test_rejects_empty_evidence(self) -> None:
        review = SourcedPaymentNotesReview(concerns=[SourcedPaymentConcern(note_index=0, reason="Anulación")])
        with patch("documents.payment_notes.create_extractor") as factory:
            factory.return_value.__enter__.return_value.run.return_value = review
            with self.assertRaisesRegex(ValueError, "evidence is absent"):
                review_payment_notes([""], {})

    def test_keeps_grounded_concern_with_conflicting_payment_terms(self) -> None:
        review = SourcedPaymentNotesReview(concerns=[SourcedPaymentConcern(note_index=0, reason="Requiere revisión")])
        with patch("documents.payment_notes.create_extractor") as factory:
            factory.return_value.__enter__.return_value.run.return_value = review
            result = review_payment_notes(["Pago a 30 días. Pedido anulado."], {})
        self.assertEqual(result.concerns[0].evidence, "Pago a 30 días. Pedido anulado.")


class PaymentPolicyTests(unittest.TestCase):
    def test_payment_restriction_escalates_but_never_overrides_paid_erp(self) -> None:
        invoice = InvoiceExtraction(
            invoice_number="INV-1", supplier_name="Proveedor", supplier_nif="B12345678",
            iban="ES123", invoice_date="2026-01-01", purchase_order="PO-2026-0001",
            line_items=[InvoiceLine(description="Servicio", amount="100")],
            tax_base="100", vat_rate="21", vat_amount="21", total="121",
            notes=["Pedido anulado."], uncertainties=[],
        )
        entry = ErpEntry(entry_id="AS-1", order_id=invoice.purchase_order, supplier_id="P1", tax_id=invoice.supplier_nif,
                         raw_amount="121,00", amount=Decimal("121"), status="PENDIENTE", raw_date="01/01/2026")
        review = PaymentNotesReview(concerns=[PaymentConcern(evidence="Pedido anulado.", reason="Anulación")])
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "rules.xlsx"
            book = openpyxl.Workbook()
            suppliers = book.create_sheet("Proveedores")
            suppliers.append(["id", "name", "nif", "iban"])
            suppliers.append(["P1", "Proveedor", invoice.supplier_nif, invoice.iban])
            orders = book.create_sheet("Pedidos_2026")
            orders.append(["id", "supplier", "nif", "total"])
            orders.append([invoice.purchase_order, "P1", invoice.supplier_nif, 121])
            book.create_sheet("pendiente_revisar")
            book.save(path)
            book.close()
            rules = PaymentRules(path, [entry], date(2026, 9, 19))
            with patch("documents.classification.review_payment_notes", return_value=review) as assess:
                decision = rules.classify(invoice)
                self.assertEqual(decision.classification, "ESCALAR")
                self.assertFalse(decision.checks["payment_notes"])
                self.assertTrue(all(passed for name, passed in decision.checks.items() if name != "payment_notes"))
                self.assertIn("Pedido anulado.", decision.reasons[0])
                assess.reset_mock()
                entry.status = "PAGADA"
                self.assertEqual(rules.classify(invoice).classification, "NO_PAGAR")
                assess.assert_not_called()


if __name__ == "__main__":
    unittest.main()
