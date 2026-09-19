import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from documents.classification import PaymentRules
from pipeline.extraction_4.features import InvoiceFeatures, InvoiceLine
from documents.payment_notes import PaymentConcern, PaymentNotesReview
from erp import ErpEntry


class ClassificationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        workbook = Path(self.directory.name) / "rules.xlsx"
        book = openpyxl.Workbook()
        suppliers = book.active
        suppliers.title = "Proveedores"
        suppliers.append(["ID", "Name", "NIF", "IBAN"])
        suppliers.append(["NEW", "Proveedor Nuevo", "B12345678", "ES00 1234"])
        orders = book.create_sheet("Pedidos_2026")
        orders.append(["Order", "Supplier", "NIF", "Amount"])
        orders.append(["PO-2026-9999", "NEW", "B12345678", 121])
        book.create_sheet("pendiente_revisar")
        book.save(workbook)
        book.close()
        self.rules = PaymentRules(workbook, [ErpEntry(
            entry_id="AS-NEW", order_id="PO-2026-9999", supplier_id="NEW", tax_id="B12345678",
            raw_amount="121,00", amount=Decimal("121.00"), status="PENDIENTE", raw_date="01/09/2026",
        )], date(2026, 9, 19))
        self.invoice = InvoiceFeatures(
            invoice_number="NEW-1", supplier_name="Proveedor Nuevo", supplier_nif="B12345678",
            iban="ES001234", invoice_date="2026-09-01", purchase_order="PO-2026-9999",
            line_items=[InvoiceLine(description="Servicio", amount="100.00")],
            tax_base="100.00", vat_rate="21", vat_amount="21.00", total="121.00",
            notes=[], uncertainties=[],
        )
        self.notes = patch("documents.classification.review_payment_notes", return_value=PaymentNotesReview(concerns=[]))
        self.review = self.notes.start()
        self.addCleanup(self.notes.stop)

    def test_new_supplier_and_order_pay_without_dataset_identifiers(self) -> None:
        self.assertEqual(self.rules.classify(self.invoice).classification, "PAGAR")

    def test_invalid_fields_escalate(self) -> None:
        cases = [
            {"supplier_nif": "B87654321"}, {"iban": "ES009999"},
            {"purchase_order": "PO-2026-9998"}, {"total": "122.00"},
            {"vat_amount": "22.00"}, {"invoice_date": "2026-09-20"},
            {"invoice_date": "2026-02-30"}, {"tax_base": ""},
            {"total": "NaN"}, {"line_items": []},
            {"uncertainties": ["Identificador ilegible"]},
        ]
        for changes in cases:
            with self.subTest(changes=changes):
                self.assertEqual(self.rules.classify(self.invoice.model_copy(update=changes)).classification, "ESCALAR")

    def test_cent_tolerance_boundary(self) -> None:
        self.assertEqual(self.rules.classify(self.invoice.model_copy(update={"total": "121.01"})).classification, "PAGAR")
        self.assertEqual(self.rules.classify(self.invoice.model_copy(update={"total": "121.02"})).classification, "ESCALAR")

    def test_pending_review_and_duplicate_orders_escalate(self) -> None:
        self.assertEqual(self.rules.classify(self.invoice, duplicate_order=True).classification, "ESCALAR")
        self.rules.pending_review.add(self.invoice.purchase_order)
        self.assertEqual(self.rules.classify(self.invoice).classification, "ESCALAR")

    def test_paid_prevents_payment_even_with_other_anomalies(self) -> None:
        self.rules.entries[0].status = "PAGADA"
        invoice = self.invoice.model_copy(update={"iban": "WRONG"})
        self.assertEqual(self.rules.classify(invoice).classification, "NO_PAGAR")
        self.review.assert_not_called()

    def test_erp_and_order_supplier_checks(self) -> None:
        self.rules.entries[0].tax_id = "B87654321"
        self.assertEqual(self.rules.classify(self.invoice).classification, "ESCALAR")
        self.rules.entries[0].tax_id = "B12345678"
        self.rules.orders = [("PO-2026-9999", "OTHER", "B12345678", 121)]
        self.assertEqual(self.rules.classify(self.invoice).classification, "ESCALAR")

    def test_commercial_hold_escalates_with_valid_arithmetic(self) -> None:
        self.review.return_value = PaymentNotesReview(concerns=[PaymentConcern(
            evidence="Factura anulada", reason="La factura está anulada.",
        )])
        invoice = self.invoice.model_copy(update={"notes": ["Factura anulada"]})
        decision = self.rules.classify(invoice)
        self.assertEqual(decision.classification, "ESCALAR")
        self.assertIn("Factura anulada", decision.reasons[0])

    def test_uses_client_normalized_amount_and_rejects_unreadable_amount(self) -> None:
        self.rules.entries[0].raw_amount = "121.00"
        self.assertEqual(self.rules.classify(self.invoice).classification, "PAGAR")
        self.rules.entries[0].amount = None
        result = self.rules.classify(self.invoice)
        self.assertEqual(result.classification, "ESCALAR")
        self.assertFalse(result.checks["erp_amount"])
