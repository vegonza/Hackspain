import re
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Literal

import openpyxl
from pydantic import BaseModel

from pipeline.extraction_4.extraction import InvoiceExtraction
from documents.payment_notes import review_payment_notes
from erp import ErpEntry

TOLERANCE = Decimal("0.01")


class Decision(BaseModel):
    classification: Literal["PAGAR", "NO_PAGAR", "ESCALAR"]
    reasons: list[str]
    checks: dict[str, bool]


def normalized(value: str) -> str:
    return re.sub(r"\s+", "", value).upper()


class PaymentRules:
    def __init__(self, workbook: Path, entries: list[ErpEntry], evaluation_date: date) -> None:
        book = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
        self.suppliers = [row for row in book["Proveedores"].iter_rows(min_row=2, values_only=True) if row[0]]
        self.orders = [row for row in book["Pedidos_2026"].iter_rows(min_row=2, values_only=True) if row[0]]
        self.pending_review = {
            match
            for row in book["pendiente_revisar"].values
            for cell in row
            for match in re.findall(r"PO-\d{4}-\d+", str(cell))
        }
        book.close()
        self.entries = entries
        self.evaluation_date = evaluation_date

    def classify(self, invoice: InvoiceExtraction, duplicate_order: bool = False) -> Decision:
        checks: dict[str, bool] = {}
        reasons = list(invoice.uncertainties)

        def check(name: str, passed: bool, reason: str) -> None:
            checks[name] = passed
            if not passed:
                reasons.append(reason)

        suppliers = list(dict.fromkeys(row for row in self.suppliers if normalized(str(row[2])) == normalized(invoice.supplier_nif)))
        check("supplier", len(suppliers) == 1, "El NIF no identifica un proveedor único en el maestro.")
        if len(suppliers) == 1:
            supplier = suppliers[0]
            check("iban", normalized(invoice.iban) == normalized(str(supplier[3])), "El IBAN no coincide con el maestro.")
        orders = [row for row in self.orders if row[0] == invoice.purchase_order]
        check("order", len(orders) == 1, "El pedido no existe o no es único en el maestro.")
        if len(orders) == 1 and len(suppliers) == 1:
            check("order_supplier", orders[0][1] == suppliers[0][0] and normalized(str(orders[0][2] or "")) == normalized(invoice.supplier_nif), "El proveedor o NIF del pedido no coincide con la factura.")
        entries = [entry for entry in self.entries if entry.order_id == invoice.purchase_order]
        check("erp_unique", len(entries) == 1, "El pedido no tiene un asiento ERP único.")
        paid = any(entry.status == "PAGADA" for entry in entries)
        if len(entries) == 1:
            entry = entries[0]
            check("erp_pending", entry.status == "PENDIENTE", "El pedido no está PENDIENTE en el ERP.")
            if len(suppliers) == 1:
                check("erp_supplier", entry.supplier_id == suppliers[0][0] and normalized(entry.tax_id) == normalized(invoice.supplier_nif), "El proveedor o NIF del ERP no coincide con la factura.")
        try:
            base, rate, vat, total = map(Decimal, (invoice.tax_base, invoice.vat_rate, invoice.vat_amount, invoice.total))
            if not all(value.is_finite() for value in (base, rate, vat, total)):
                raise InvalidOperation
            check("vat", abs(vat - (base * rate / 100).quantize(TOLERANCE, rounding=ROUND_HALF_UP)) <= TOLERANCE, "El IVA no corresponde a la base y al tipo impreso.")
            check("total", abs(total - base - vat) <= TOLERANCE, "El total no coincide con base más IVA.")
            line_sum = sum((Decimal(line.amount) for line in invoice.line_items), Decimal(0))
            check("line_sum", bool(invoice.line_items) and abs(line_sum - base) <= TOLERANCE, "La suma de conceptos no coincide con la base.")
            if len(orders) == 1:
                check("order_amount", abs(total - Decimal(str(orders[0][3]))) <= TOLERANCE, "El total no coincide con el importe del pedido.")
            if len(entries) == 1:
                amount = entries[0].amount
                check("erp_amount", amount is not None and abs(total - amount) <= TOLERANCE, "El importe ERP no es legible o no coincide con el total.")
        except InvalidOperation:
            check("amounts_readable", False, "Faltan importes legibles o no son decimales válidos.")
        try:
            invoice_date = date.fromisoformat(invoice.invoice_date)
            check("date", invoice_date <= self.evaluation_date, "La fecha de factura es futura.")
        except ValueError:
            check("date", False, "La fecha de factura no es válida o no se puede interpretar.")
        check("pending_review", invoice.purchase_order not in self.pending_review, "El pedido tiene una revisión pendiente en el Excel.")
        check("duplicate_order", not duplicate_order, "Varias facturas reclaman el mismo pedido.")
        if paid:
            return Decision(classification="NO_PAGAR", reasons=["El ERP registra el pedido como PAGADA; nunca pagar dos veces.", *reasons], checks=checks)
        review = review_payment_notes(invoice.notes, checks)
        checks["payment_notes"] = not review.concerns
        reasons.extend(f"{concern.reason} Evidencia: {concern.evidence}" for concern in review.concerns)
        return Decision(classification="ESCALAR" if reasons else "PAGAR", reasons=reasons or ["Identidad, pedido, ERP, importes y fecha verificados."], checks=checks)
